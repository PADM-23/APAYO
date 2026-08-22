#!/usr/bin/env python3
"""경북 병원·약국 기본정보와 진료과목을 앱용 JSON으로 변환한다."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import sys
import tempfile
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import load_workbook
except ImportError:
    print(
        "openpyxl이 필요합니다. 다음 명령을 먼저 실행하세요:\n"
        "python3 -m pip install -r DataPipeline/requirements.txt",
        file=sys.stderr,
    )
    raise SystemExit(1)


PIPELINE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = PIPELINE_DIR.parent
DEFAULT_HOSPITALS = PIPELINE_DIR / "raw" / "hospitals_2026-06.xlsx"
DEFAULT_PHARMACIES = PIPELINE_DIR / "raw" / "pharmacies_2026-06.xlsx"
DEFAULT_DEPARTMENTS = PIPELINE_DIR / "raw" / "departments_2026-06.xlsx"
DEFAULT_OUTPUT = PIPELINE_DIR / "output" / "GyeongbukMedicalFacilities.json"
DEFAULT_APP_OUTPUT = (
    PROJECT_DIR / "APAYO" / "Resources" / "Data" / "GyeongbukMedicalFacilities.json"
)

FACILITY_REQUIRED_COLUMNS = {
    "암호화요양기호",
    "요양기관명",
    "종별코드명",
    "시도코드명",
    "시군구코드명",
    "주소",
    "전화번호",
    "좌표(X)",
    "좌표(Y)",
}

DEPARTMENT_REQUIRED_COLUMNS = {
    "암호화요양기호",
    "진료과목코드",
    "진료과목코드명",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="심평원 XLSX 세 개를 결합해 경북 병원·약국 JSON을 생성합니다."
    )
    parser.add_argument("--hospitals", type=Path, default=DEFAULT_HOSPITALS)
    parser.add_argument("--pharmacies", type=Path, default=DEFAULT_PHARMACIES)
    parser.add_argument("--departments", type=Path, default=DEFAULT_DEPARTMENTS)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--app-output", type=Path, default=DEFAULT_APP_OUTPUT)
    return parser.parse_args()


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def make_id(encrypted_facility_id: str) -> str:
    return hashlib.sha256(encrypted_facility_id.encode("utf-8")).hexdigest()[:16]


def validate_columns(headers: Iterable[Any], required: set[str], file_label: str) -> dict[str, int]:
    normalized = [clean_text(value) for value in headers]
    index = {name: position for position, name in enumerate(normalized) if name}
    missing = sorted(required - set(index))
    if missing:
        raise ValueError(
            f"{file_label} 파일의 형식이 예상과 다릅니다. 없는 컬럼: {', '.join(missing)}"
        )
    return index


def load_gyeongbuk_facilities(
    path: Path, *, category: str, file_label: str
) -> dict[str, dict[str, Any]]:
    if not path.is_file():
        raise FileNotFoundError(f"{file_label} 파일이 없습니다: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    index = validate_columns(next(rows), FACILITY_REQUIRED_COLUMNS, file_label)
    facilities: dict[str, dict[str, Any]] = {}

    try:
        for row in rows:
            if clean_text(row[index["시도코드명"]]) != "경북":
                continue

            encrypted_id = clean_text(row[index["암호화요양기호"]])
            name = clean_text(row[index["요양기관명"]])
            if not encrypted_id or not name:
                continue
            if encrypted_id in facilities:
                raise ValueError(f"{file_label}에 암호화요양기호가 중복되었습니다: {name}")

            longitude = row[index["좌표(X)"]]
            latitude = row[index["좌표(Y)"]]
            facilities[encrypted_id] = {
                "id": make_id(encrypted_id),
                "category": category,
                "name": name,
                "type": clean_text(row[index["종별코드명"]]) or "기타",
                "district": clean_text(row[index["시군구코드명"]]) or "",
                "address": clean_text(row[index["주소"]]) or "",
                "phone": clean_text(row[index["전화번호"]]),
                "latitude": float(latitude) if latitude is not None else None,
                "longitude": float(longitude) if longitude is not None else None,
                "departments": [],
            }
    finally:
        workbook.close()

    return facilities


def attach_departments(path: Path, facilities: dict[str, dict[str, Any]]) -> int:
    if not path.is_file():
        raise FileNotFoundError(f"진료과목 파일이 없습니다: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    index = validate_columns(next(rows), DEPARTMENT_REQUIRED_COLUMNS, "진료과목")
    department_names: dict[str, set[str]] = defaultdict(set)
    joined_rows = 0

    try:
        for row in rows:
            encrypted_id = clean_text(row[index["암호화요양기호"]])
            if not encrypted_id or encrypted_id not in facilities:
                continue

            department_name = clean_text(row[index["진료과목코드명"]])
            if department_name:
                department_names[encrypted_id].add(department_name)
                joined_rows += 1
    finally:
        workbook.close()

    for encrypted_id, facility in facilities.items():
        if facility["category"] == "hospital":
            facility["departments"] = sorted(
                department_names.get(encrypted_id, set())
            )

    return joined_rows


def validate_result(facilities: list[dict[str, Any]]) -> dict[str, int]:
    ids = [facility["id"] for facility in facilities]
    if len(ids) != len(set(ids)):
        raise ValueError("생성된 의료시설 ID가 중복되었습니다.")

    invalid_coordinates = 0
    missing_coordinates = 0
    missing_phone = 0
    missing_departments = 0

    hospital_count = 0
    pharmacy_count = 0

    for facility in facilities:
        if facility["category"] == "hospital":
            hospital_count += 1
        elif facility["category"] == "pharmacy":
            pharmacy_count += 1
        else:
            raise ValueError(f"알 수 없는 의료시설 category: {facility['category']}")

        latitude = facility["latitude"]
        longitude = facility["longitude"]
        if latitude is None or longitude is None:
            missing_coordinates += 1
        elif not (
            math.isfinite(latitude)
            and math.isfinite(longitude)
            and 33 <= latitude <= 39
            and 124 <= longitude <= 132
        ):
            invalid_coordinates += 1

        if facility["phone"] is None:
            missing_phone += 1
        if facility["category"] == "hospital" and not facility["departments"]:
            missing_departments += 1

    if missing_coordinates:
        raise ValueError(f"좌표가 없는 의료시설이 {missing_coordinates}개 있습니다.")
    if invalid_coordinates:
        raise ValueError(f"대한민국 범위를 벗어난 좌표가 {invalid_coordinates}개 있습니다.")

    return {
        "facility_count": len(facilities),
        "hospital_count": hospital_count,
        "pharmacy_count": pharmacy_count,
        "missing_coordinates": missing_coordinates,
        "missing_phone": missing_phone,
        "missing_departments": missing_departments,
    }


def write_json_atomically(path: Path, facilities: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(file_descriptor, "w", encoding="utf-8") as file:
            json.dump(facilities, file, ensure_ascii=False, separators=(",", ":"))
            file.write("\n")
        os.replace(temporary_name, path)
        path.chmod(0o644)
    except Exception:
        try:
            os.unlink(temporary_name)
        except FileNotFoundError:
            pass
        raise


def main() -> int:
    args = parse_args()

    print("1/5 경북 병원 기본정보를 읽는 중...")
    facility_map = load_gyeongbuk_facilities(
        args.hospitals, category="hospital", file_label="병원 기본정보"
    )

    print("2/5 진료과목을 암호화요양기호로 연결하는 중...")
    joined_department_rows = attach_departments(args.departments, facility_map)

    print("3/5 경북 약국 기본정보를 읽는 중...")
    pharmacies = load_gyeongbuk_facilities(
        args.pharmacies, category="pharmacy", file_label="약국 기본정보"
    )
    duplicate_keys = set(facility_map) & set(pharmacies)
    if duplicate_keys:
        raise ValueError(
            f"병원과 약국 파일에 중복된 암호화요양기호가 {len(duplicate_keys)}개 있습니다."
        )
    facility_map.update(pharmacies)

    facilities = sorted(
        facility_map.values(),
        key=lambda facility: (facility["category"], facility["name"], facility["id"]),
    )

    print("4/5 결과를 검증하는 중...")
    stats = validate_result(facilities)

    print("5/5 JSON 파일을 저장하는 중...")
    write_json_atomically(args.output, facilities)
    if args.app_output.resolve() != args.output.resolve():
        write_json_atomically(args.app_output, facilities)

    print("완료")
    print(f"- 경북 의료시설 전체: {stats['facility_count']:,}개")
    print(f"- 병원·의원: {stats['hospital_count']:,}개")
    print(f"- 약국: {stats['pharmacy_count']:,}개")
    print(f"- 연결된 진료과목 행: {joined_department_rows:,}개")
    print(f"- 진료과목 없음: {stats['missing_departments']:,}개")
    print(f"- 전화번호 없음: {stats['missing_phone']:,}개")
    print(f"- 좌표 없음: {stats['missing_coordinates']:,}개")
    print(f"- 확인용 JSON: {args.output}")
    print(f"- 앱 번들 JSON: {args.app_output}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (FileNotFoundError, ValueError) as error:
        print(f"오류: {error}", file=sys.stderr)
        raise SystemExit(1)
